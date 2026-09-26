"""notes/part*.md 를 합쳐 해설.md 와 해설.docx 를 만든다."""
import glob, os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TITLE = "알바로 이야기 (85-110) — 문장별 스페인어 해설"

parts = sorted(glob.glob(os.path.join(HERE, "part*.md")))
md = f"# {TITLE}\n\n" + "\n\n".join(open(p, encoding="utf-8").read().strip() for p in parts) + "\n"
open(os.path.join(ROOT, "해설.md"), "w", encoding="utf-8").write(md)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Malgun Gothic"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
def add_runs(par, text):
    for t in TOKEN.split(text):
        if not t: continue
        if t.startswith("**") and t.endswith("**"):
            par.add_run(t[2:-2]).bold = True
        elif t.startswith("`") and t.endswith("`"):
            r = par.add_run(t[1:-1]); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9A); r.bold = True
        else:
            par.add_run(t)

for line in md.splitlines():
    s = line.rstrip()
    if not s.strip() or s.strip() == "---": continue
    if s.startswith("# "): doc.add_heading(s[2:], 0)
    elif s.startswith("## "): doc.add_heading(s[3:], 1)
    elif s.startswith("### "): add_runs(doc.add_heading("", 3), s[4:])
    elif s.startswith("> "):
        p = doc.add_paragraph(style="Intense Quote"); add_runs(p, s[2:])
    elif re.match(r"^\s*[-*] ", s):
        lvl = (len(s) - len(s.lstrip())) // 2
        p = doc.add_paragraph(style="List Bullet 2" if lvl else "List Bullet")
        add_runs(p, re.sub(r"^\s*[-*] ", "", s))
    else:
        add_runs(doc.add_paragraph(), s)

doc.save(os.path.join(ROOT, "해설.docx"))
print(len(parts), "parts,", md.count("\n### "), "sentences")
